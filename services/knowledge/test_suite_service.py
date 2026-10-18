"""Project-scoped QA test suite ingestion; never fabricates execution results."""
from __future__ import annotations
import csv
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any
from sqlalchemy.orm import Session
from db_models import KnowledgeDocument

SUPPORTED = {'.xlsx', '.csv', '.json'}
MAX_BYTES = 12 * 1024 * 1024


def _norm(value: Any) -> str:
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def _parse_rows(filename: str, data: bytes):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Test suite is empty or exceeds 12 MB')
    ext = Path(filename).suffix.lower()
    if ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in book:
                iterator = sheet.iter_rows(values_only=True)
                header = next(iterator, ())
                if not header:
                    continue
                fields = [str(v or '').strip() for v in header]
                for row_number, values in enumerate(iterator, 2):
                    if any(v is not None and str(v).strip() for v in values):
                        yield sheet.title, row_number, dict(zip(fields, values))
        finally:
            book.close()
    elif ext == '.csv':
        reader = csv.DictReader(io.StringIO(data.decode('utf-8-sig')))
        for row_number, row in enumerate(reader, 2):
            yield 'CSV', row_number, row
    elif ext == '.json':
        obj = json.loads(data.decode('utf-8-sig'))
        records = obj if isinstance(obj, list) else obj.get('test_cases', obj.get('tests', []))
        if not isinstance(records, list):
            raise ValueError('JSON test suite requires a list of test cases')
        for row_number, record in enumerate(records, 1):
            if isinstance(record, dict):
                yield 'JSON', row_number, record
    else:
        raise ValueError('Supported test suites: XLSX, CSV, JSON')


ALIASES = {
    'test_id': ('testid', 'testcaseid', 'caseid', 'id'),
    'scenario': ('scenario', 'testscenario', 'testcase', 'description', 'testdescription', 'testname', 'teststeps', 'testcase description'),
    'input': ('input', 'inputdata', 'inputcountrycode', 'sampledata', 'sampleinput', 'testdata', 'testinput', 'request', 'requestjson', 'countrycode'),
    'expected': ('expected', 'expectedresult', 'expectedoutcome', 'expectedoutput', 'expectedresponse', 'expectedbehavior', 'expectedresults'),
    'jira': ('jira', 'jiraid', 'jirakey', 'issuekey', 'jiraissue', 'jiraissuekey'),
    'attribute': ('attribute', 'field', 'fieldname', 'targetattribute', 'javaattribute', 'javafield'),
    'operation': ('operation', 'businessoperation', 'endpoint', 'method', 'apioperation'),
}


def parse_test_cases(filename: str, data: bytes, jira_key: str | None = None) -> list[dict]:
    cases = []
    for sheet, row_number, raw in _parse_rows(filename, data):
        lookup = {_norm(k): v for k, v in raw.items() if k}
        def field(name):
            return next((lookup[k] for k in ALIASES[name] if k in lookup and lookup[k] is not None), None)
        test_id = field('test_id')
        # Ignore overview, mapping verification and other non-test sheets.
        if not test_id or not (field('scenario') or field('expected')):
            continue
        jira = jira_key or field('jira') or re.search(r'\b[A-Z][A-Z0-9]+-\d+\b', filename, re.I)
        if hasattr(jira, 'group'):
            jira = jira.group(0)
        cases.append({
            'test_id': str(test_id).strip(), 'scenario': str(field('scenario') or '').strip(),
            'input': str(field('input') or '').strip(), 'expected': str(field('expected') or '').strip(),
            'jira': str(jira or '').upper(), 'attribute': str(field('attribute') or '').strip().split('.')[-1],
            'operation': str(field('operation') or '').strip(), 'execution_status': 'NOT_VERIFIED',
            'sheet': sheet, 'row': row_number,
        })
    return cases


class TestSuiteService:
    def ingest(self, db: Session, *, filename: str, data: bytes, project_path: str,
               source_ref: str, jira_key: str | None = None, source_kind: str = 'FOLDER',
               source_id: int | None = None) -> dict:
        cases = parse_test_cases(filename, data, jira_key)
        if not cases:
            return {'status': 'NO_TEST_CASES', 'cases': 0, 'filename': filename}
        digest = hashlib.sha256(data).hexdigest()
        old = db.query(KnowledgeDocument).filter(
            KnowledgeDocument.project_path == project_path,
            KnowledgeDocument.source_ref == source_ref,
            KnowledgeDocument.source_type == 'TEST',
            KnowledgeDocument.status == 'ACTIVE').all()
        if any(json.loads(row.metadata_json or '{}').get('sha256') == digest for row in old):
            return {'status': 'UNCHANGED', 'cases': len(cases), 'filename': filename}
        for row in old:
            row.status = 'SUPERSEDED'
        content = '\n'.join(f"Test {c['test_id']} | JIRA {c['jira']} | Attribute {c['attribute']} | Operation {c['operation']} | Scenario {c['scenario']} | Input {c['input']} | Expected {c['expected']} | Execution NOT_VERIFIED" for c in cases)
        doc = KnowledgeDocument(source_type='TEST', title=filename, content=content,
            project_path=project_path, source_ref=source_ref, status='ACTIVE',
            metadata_json=json.dumps({'sha256': digest, 'source_kind': source_kind,
                'jira': jira_key, 'source_id': source_id, 'cases': cases}, ensure_ascii=False, default=str))
        db.add(doc)
        db.commit()
        return {'status': 'IMPORTED', 'cases': len(cases), 'filename': filename, 'document_id': doc.id}

    def search(self, db: Session, *, project_path: str, jira: str = '', attribute: str = '', limit: int = 100) -> dict:
        # IntelliJ and backend may supply the same Windows path with different case or separators.
        # Keep retrieval project-scoped; never silently include another project.
        def scope(value: str) -> str:
            return re.sub(r'/+', '/', str(value or '').strip().replace('\\', '/')).rstrip('/').casefold()
        rows = db.query(KnowledgeDocument).filter(
            KnowledgeDocument.source_type == 'TEST', KnowledgeDocument.status == 'ACTIVE').all()
        rows = [row for row in rows if scope(row.project_path) == scope(project_path)]
        found = []
        for doc in rows:
            meta = json.loads(doc.metadata_json or '{}')
            for case in meta.get('cases', []):
                if jira and jira.upper() != str(case.get('jira') or '').upper():
                    continue
                if attribute and _norm(attribute) not in _norm(' '.join(str(case.get(k) or '') for k in ('attribute', 'scenario', 'input', 'expected'))):
                    continue
                found.append({**case, 'source_filename': doc.title, 'source_ref': doc.source_ref,
                    'source_kind': meta.get('source_kind'), 'source_jira': meta.get('jira'),
                    'source_id': meta.get('source_id'), 'document_id': doc.id})
        return {'project_path': project_path, 'count': len(found), 'test_cases': found[:max(1, min(limit, 500))]}
