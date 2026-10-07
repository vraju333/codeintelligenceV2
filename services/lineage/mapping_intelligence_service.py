from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import or_, inspect, text
from sqlalchemy.orm import Session

from config import settings
from db_models import MappingDefinition, MappingDocument


class MappingIntelligenceService:
    """Phase 6 source-agnostic mapping-document intelligence.

    Excel is the first-class ingestion format. Every imported row preserves its
    workbook/sheet/row provenance so answers can always point back to evidence.
    """

    HEADER_ALIASES = {
        "source_type": {"source type", "input type", "format", "source format", "input format"},
        "source_system": {"source system", "input system", "source application", "source app"},
        "source_path": {"input attribute", "input path", "source", "source field", "source attribute", "source path", "xml node", "xpath", "json node", "json path", "db column", "source column"},
        "mapping_rule": {"mapping", "mapping rule", "transformation", "transform", "logic", "rule", "conversion"},
        "target_attribute": {"java attribute", "java attribute to be mapped to", "target", "target field", "target attribute", "java field", "application attribute"},
        "null_rule": {"if null", "null handling", "null rule", "default", "default value", "null/default handling"},
        "validation_rule": {"validation", "validation rule", "validations", "constraint"},
        "comments": {"comments", "comment", "remarks", "notes", "description", "business comments"},
    }

    def _project_path(self) -> str:
        value = getattr(settings, "PYTHON_PROJECT_PATH", None) or getattr(settings, "JAVA_PROJECT_PATH", None) or "ACTIVE_PROJECT"
        try:
            return str(Path(str(value)).expanduser().resolve())
        except Exception:
            return str(value)

    @staticmethod
    def _norm_header(value: Any) -> str:
        text = str(value or "").strip().lower().replace("_", " ").replace("-", " ")
        return re.sub(r"\s+", " ", text)

    def _canonical_header(self, value: Any) -> str | None:
        normalized = self._norm_header(value)
        for canonical, aliases in self.HEADER_ALIASES.items():
            if normalized in aliases:
                return canonical
        return None

    @staticmethod
    def _clean(value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text if text else None

    @staticmethod
    def _infer_source_type(path: str | None, explicit: str | None) -> str:
        if explicit:
            return explicit.strip().upper()
        value = (path or "").strip()
        lower = value.lower()
        if value.startswith("/") or lower.startswith("xpath:") or "</" in lower:
            return "XML"
        if lower.startswith(("$.", "json:", "jsonpath:")):
            return "JSON"
        if re.match(r"^[A-Za-z0-9_]+\.[A-Z0-9_]+$", value) and value.upper() == value:
            return "DB"
        if lower.startswith(("db:", "table:", "column:")):
            return "DB"
        if lower.startswith(("kafka:", "topic:")):
            return "KAFKA"
        return "UNKNOWN"

    @staticmethod
    def _target_parts(target: str | None) -> tuple[str | None, str | None]:
        value = str(target or "").strip()
        if not value:
            return None, None
        value = value.replace("#", ".")
        if "." in value:
            owner, attribute = value.rsplit(".", 1)
            return owner.strip() or None, attribute.strip() or None
        return None, value

    def import_excel(self, db: Session, *, filename: str, content: bytes, title: str | None = None,
                     document_version: str | None = None, mapping_family: str | None = None, source_ref: str | None = None) -> dict:
        if not filename.lower().endswith((".xlsx", ".xlsm")):
            raise ValueError("Phase 6 mapping upload currently accepts .xlsx or .xlsm files")
        if not content:
            raise ValueError("mapping workbook is empty")

        self._ensure_mapping_family_column(db)
        project_path = self._project_path()
        checksum = hashlib.sha256(content).hexdigest()
        existing = db.query(MappingDocument).filter(
            MappingDocument.project_path == project_path,
            MappingDocument.checksum_sha256 == checksum,
            MappingDocument.status == "ACTIVE",
        ).first()
        if existing:
            return {"status": "ALREADY_IMPORTED", "document": self._document_dict(existing), "rows_imported": existing.row_count}

        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        document = MappingDocument(
            project_path=project_path,
            filename=filename,
            title=(title or Path(filename).stem),
            mapping_family=(mapping_family or title or Path(filename).stem).strip(),
            document_version=(document_version or None),
            source_ref=(source_ref or filename),
            checksum_sha256=checksum,
            status="ACTIVE",
        )
        db.add(document)
        db.flush()

        imported: list[MappingDefinition] = []
        skipped: list[dict] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header_index, header_map = self._find_header(rows)
            if header_index is None:
                skipped.append({"sheet": sheet.title, "reason": "No recognizable mapping header row"})
                continue
            for row_number, values in enumerate(rows[header_index + 1:], start=header_index + 2):
                data = {key: self._clean(values[index]) if index < len(values) else None for key, index in header_map.items()}
                source_path = data.get("source_path")
                target = data.get("target_attribute")
                if not source_path and not target:
                    continue
                if not source_path or not target:
                    skipped.append({"sheet": sheet.title, "row": row_number, "reason": "Source or target is missing"})
                    continue
                owner, attribute = self._target_parts(target)
                mapping = MappingDefinition(
                    document_id=document.id,
                    project_path=project_path,
                    sheet_name=sheet.title,
                    row_number=row_number,
                    source_type=self._infer_source_type(source_path, data.get("source_type")),
                    source_system=data.get("source_system"),
                    source_path=source_path,
                    mapping_rule=data.get("mapping_rule"),
                    target_class=owner,
                    target_attribute=attribute,
                    target_expression=target,
                    null_rule=data.get("null_rule"),
                    validation_rule=data.get("validation_rule"),
                    comments=data.get("comments"),
                    raw_row_json=json.dumps({str(rows[header_index][i] or f"Column {i+1}"): values[i] if i < len(values) else None for i in range(len(rows[header_index]))}, default=str),
                    status="ACTIVE",
                )
                db.add(mapping)
                imported.append(mapping)

        document.row_count = len(imported)
        document.sheet_count = len(workbook.sheetnames)
        db.commit()
        db.refresh(document)
        for row in imported:
            db.refresh(row)

        graph = self.sync_document_to_neo4j(db, document.id)
        rag = self._index_mapping_document(db, document, imported)
        return {
            "status": "IMPORTED",
            "document": self._document_dict(document),
            "rows_imported": len(imported),
            "skipped": skipped,
            "neo4j": graph,
            "rag": rag,
        }

    def _index_mapping_document(self, db: Session, document: MappingDocument, rows: list[MappingDefinition]) -> dict:
        """Make mapping definitions semantically discoverable without making RAG authoritative."""
        try:
            from services.knowledge.knowledge_ingestion_service import KnowledgeIngestionService
            lines = [f"Mapping document: {document.filename}", f"Version: {document.document_version or 'unspecified'}"]
            for row in rows:
                lines.append(
                    f"Sheet {row.sheet_name}, row {row.row_number}: {row.source_type} {row.source_path} -> "
                    f"{row.target_expression}; mapping={row.mapping_rule or ''}; null={row.null_rule or ''}; "
                    f"validation={row.validation_rule or ''}; comments={row.comments or ''}"
                )
            result = KnowledgeIngestionService().ingest(
                db, source_type="MAPPING", title=document.title, content="\n".join(lines),
                source_ref=document.source_ref or document.filename,
                metadata={"mapping_document_id": document.id, "filename": document.filename, "document_version": document.document_version},
            )
            return {"status": "INDEXED", "knowledge_document_id": (result.get("document") or {}).get("id")}
        except Exception as exc:
            # Mapping rows remain authoritative even when semantic indexing is unavailable.
            return {"status": "UNAVAILABLE", "error": str(exc)}

    def _find_header(self, rows: list[tuple]) -> tuple[int | None, dict[str, int]]:
        best_index, best_map = None, {}
        for index, row in enumerate(rows[:25]):
            found: dict[str, int] = {}
            for col, value in enumerate(row):
                canonical = self._canonical_header(value)
                if canonical and canonical not in found:
                    found[canonical] = col
            if "source_path" in found and "target_attribute" in found and len(found) > len(best_map):
                best_index, best_map = index, found
        return best_index, best_map

    def list_documents(self, db: Session) -> dict:
        self._ensure_mapping_family_column(db)
        rows = db.query(MappingDocument).filter(
            MappingDocument.project_path == self._project_path(), MappingDocument.status == "ACTIVE"
        ).order_by(MappingDocument.id.desc()).all()
        return {"documents": [self._document_dict(row) for row in rows]}

    def document_rows(self, db: Session, document_id: int) -> dict:
        document = db.query(MappingDocument).filter(MappingDocument.id == document_id).first()
        if not document:
            raise ValueError("mapping document not found")
        rows = db.query(MappingDefinition).filter(
            MappingDefinition.document_id == document_id, MappingDefinition.status == "ACTIVE"
        ).order_by(MappingDefinition.sheet_name, MappingDefinition.row_number).all()
        return {"document": self._document_dict(document), "mappings": [self._mapping_dict(row, document) for row in rows]}

    def search(self, db: Session, query: str, limit: int = 50) -> dict:
        value = str(query or "").strip()
        if not value:
            raise ValueError("query is required")
        pattern = f"%{value}%"
        rows = db.query(MappingDefinition).filter(
            MappingDefinition.project_path == self._project_path(), MappingDefinition.status == "ACTIVE",
            or_(
                MappingDefinition.source_path.ilike(pattern), MappingDefinition.target_attribute.ilike(pattern),
                MappingDefinition.target_expression.ilike(pattern), MappingDefinition.mapping_rule.ilike(pattern),
                MappingDefinition.null_rule.ilike(pattern), MappingDefinition.comments.ilike(pattern),
            )
        ).order_by(MappingDefinition.document_id.desc(), MappingDefinition.row_number).limit(max(1, min(limit, 250))).all()
        docs = {d.id: d for d in db.query(MappingDocument).filter(MappingDocument.id.in_({r.document_id for r in rows})).all()} if rows else {}
        return {"query": value, "matches": [self._mapping_dict(row, docs.get(row.document_id)) for row in rows], "count": len(rows)}

    def lineage(self, db: Session, attribute: str) -> dict:
        result = self.search(db, attribute, 250)
        exact = []
        needle = attribute.strip().lower()
        for row in result["matches"]:
            candidates = [row.get("target_attribute"), row.get("target_expression"), row.get("source_path")]
            if any(needle == str(x or "").lower() or needle in str(x or "").lower() for x in candidates):
                exact.append(row)
        conflicts = self._conflicts(exact)
        return {"attribute": attribute, "mappings": exact, "mapping_count": len(exact), "conflicts": conflicts,
                "evidence_rule": "Every mapping includes document, sheet and row provenance."}


    def graph_lineage(self, attribute: str, project: str | None = None) -> dict:
        """Traverse Phase-6 MAPS_TO evidence into the Phase-5 engineering graph.

        This is intentionally Neo4j-backed (unlike lineage(), whose source of truth
        is PostgreSQL/SQLite). It proves that an imported mapping source is joined
        to the real Java Attribute and then follows the focused Phase-5 impact path.
        """
        value = str(attribute or "").strip()
        if not value:
            raise ValueError("attribute is required")
        if not getattr(settings, "NEO4J_ENABLED", False):
            return {"status": "DISABLED", "attribute": value, "mappings": [], "engineering_impact": None}

        requested_project = str(project or "").strip()
        driver = None
        try:
            from neo4j import GraphDatabase
            driver = GraphDatabase.driver(
                settings.NEO4J_URI,
                auth=(settings.NEO4J_USERNAME, settings.NEO4J_PASSWORD),
            )
            with driver.session(database=settings.NEO4J_DATABASE) as session:
                records = session.run(
                    """
                    MATCH (s:EngineeringKnowledge:MappingSource)-[r:MAPS_TO]->(a:EngineeringKnowledge:Attribute)
                    WHERE (toLower(coalesce(a.attribute,'')) = toLower($attribute)
                           OR toLower(a.name) = toLower($attribute)
                           OR toLower(a.name) ENDS WITH '.' + toLower($attribute))
                      AND ($project = '' OR a.project = $project OR s.project = $project)
                    OPTIONAL MATCH (m:EngineeringKnowledge:MappingDefinition)-[:TARGETS]->(a)
                    WHERE m.mapping_id = r.mapping_id
                    OPTIONAL MATCH (m)-[:DEFINED_IN]->(d:EngineeringKnowledge:MappingDocument)
                    RETURN s, r, a, m, d
                    ORDER BY coalesce(d.mapping_document_id, 0), coalesce(r.row, 0)
                    """,
                    attribute=value,
                    project=requested_project,
                )
                mappings = []
                projects = set()
                for record in records:
                    source = dict(record["s"]) if record["s"] is not None else {}
                    rel = dict(record["r"]) if record["r"] is not None else {}
                    target = dict(record["a"]) if record["a"] is not None else {}
                    definition = dict(record["m"]) if record["m"] is not None else {}
                    document = dict(record["d"]) if record["d"] is not None else {}
                    if target.get("project"):
                        projects.add(str(target["project"]))
                    mappings.append({
                        "source": {
                            "type": source.get("source_type"),
                            "system": source.get("source_system"),
                            "path": source.get("name"),
                        },
                        "relationship": "MAPS_TO",
                        "target": {
                            "class": target.get("owner"),
                            "attribute": target.get("attribute") or target.get("name"),
                            "name": target.get("name"),
                            "project": target.get("project"),
                        },
                        "mapping": {
                            "mapping_id": rel.get("mapping_id"),
                            "rule": rel.get("mapping_rule") or definition.get("mapping_rule"),
                            "null_rule": rel.get("null_rule") or definition.get("null_rule"),
                            "validation_rule": rel.get("validation_rule") or definition.get("validation_rule"),
                        },
                        "provenance": {
                            "document_id": document.get("mapping_document_id"),
                            "document": rel.get("document") or document.get("filename"),
                            "document_version": document.get("version"),
                            "sheet": rel.get("sheet") or definition.get("sheet"),
                            "row": rel.get("row") or definition.get("row"),
                            "source_ref": document.get("source_ref"),
                        },
                    })

            if not mappings:
                return {
                    "status": "NO_MATCH",
                    "attribute": value,
                    "project": requested_project or None,
                    "mappings": [],
                    "mapping_count": 0,
                    "engineering_impact": None,
                }

            # Reuse the already-hardened Phase-5 semantic traversal instead of
            # duplicating its CALLS/INVOKES/COVERED_BY logic here.
            from services.cross_project.engineering_knowledge_graph_service import EngineeringKnowledgeGraphService
            impact_project = requested_project or (next(iter(projects)) if len(projects) == 1 else None)
            impact = EngineeringKnowledgeGraphService().impact(value, impact_project)
            return {
                "status": "FOUND",
                "attribute": value,
                "project": impact_project,
                "mappings": mappings,
                "mapping_count": len(mappings),
                "engineering_impact": impact,
                "evidence_rule": "MAPS_TO evidence is read from Neo4j and preserves document, sheet and row provenance.",
            }
        finally:
            if driver:
                driver.close()

    def _ensure_mapping_family_column(self, db: Session) -> None:
        """Small backward-compatible migration for existing Phase-6 databases."""
        engine = db.get_bind()
        columns = {c["name"] for c in inspect(engine).get_columns("mapping_documents")}
        if "mapping_family" not in columns:
            dialect = engine.dialect.name
            ddl = "ALTER TABLE mapping_documents ADD COLUMN mapping_family VARCHAR(500)"
            db.execute(text(ddl))
            db.commit()
        db.execute(text("UPDATE mapping_documents SET mapping_family = title WHERE mapping_family IS NULL OR TRIM(mapping_family) = ''"))
        db.commit()

    @staticmethod
    def _version_key(value: str | None) -> tuple:
        text_value = str(value or "").strip()
        parts = re.split(r"(\d+)", text_value.lower())
        return tuple(int(p) if p.isdigit() else p for p in parts)

    def list_families(self, db: Session) -> dict:
        self._ensure_mapping_family_column(db)
        docs = db.query(MappingDocument).filter(
            MappingDocument.project_path == self._project_path(), MappingDocument.status == "ACTIVE"
        ).all()
        grouped: dict[str, list[MappingDocument]] = {}
        for doc in docs:
            grouped.setdefault(doc.mapping_family or doc.title, []).append(doc)
        families = []
        for family, items in sorted(grouped.items()):
            items.sort(key=lambda d: (self._version_key(d.document_version), d.created_at or datetime.min))
            families.append({
                "mapping_family": family,
                "versions": [d.document_version for d in items],
                "document_ids": [d.id for d in items],
                "latest_version": items[-1].document_version if items else None,
                "document_count": len(items),
            })
        return {"families": families}

    def _family_documents(self, db: Session, family: str) -> list[MappingDocument]:
        self._ensure_mapping_family_column(db)
        docs = db.query(MappingDocument).filter(
            MappingDocument.project_path == self._project_path(),
            MappingDocument.status == "ACTIVE",
            MappingDocument.mapping_family.ilike(family),
        ).all()
        docs.sort(key=lambda d: (self._version_key(d.document_version), d.created_at or datetime.min))
        return docs

    @staticmethod
    def _row_signature(row: MappingDefinition) -> dict:
        return {
            "source_type": row.source_type,
            "source_system": row.source_system,
            "source_path": row.source_path,
            "target_class": row.target_class,
            "target_attribute": row.target_attribute,
            "target_expression": row.target_expression,
            "mapping_rule": row.mapping_rule,
            "null_rule": row.null_rule,
            "validation_rule": row.validation_rule,
            "comments": row.comments,
        }

    @staticmethod
    def _changes(before: dict, after: dict) -> list[dict]:
        labels = {
            "source_type": "SOURCE_TYPE", "source_system": "SOURCE_SYSTEM", "source_path": "SOURCE_PATH",
            "target_class": "TARGET_CLASS", "target_attribute": "TARGET_ATTRIBUTE", "target_expression": "TARGET",
            "mapping_rule": "MAPPING_RULE", "null_rule": "NULL_RULE", "validation_rule": "VALIDATION_RULE",
            "comments": "COMMENTS",
        }
        changes = []
        for field, change_type in labels.items():
            old, new = before.get(field), after.get(field)
            if (old or None) != (new or None):
                changes.append({"type": change_type, "field": field, "from": old, "to": new})
        return changes

    def mapping_history(self, db: Session, attribute: str, mapping_family: str | None = None) -> dict:
        needle = str(attribute or "").strip().lower()
        if not needle:
            raise ValueError("attribute is required")
        self._ensure_mapping_family_column(db)
        query = db.query(MappingDocument).filter(
            MappingDocument.project_path == self._project_path(), MappingDocument.status == "ACTIVE"
        )
        if mapping_family:
            query = query.filter(MappingDocument.mapping_family.ilike(mapping_family))
        docs = query.all()
        docs.sort(key=lambda d: ((d.mapping_family or d.title).lower(), self._version_key(d.document_version), d.created_at or datetime.min))
        timelines: dict[str, list[dict]] = {}
        for doc in docs:
            rows = db.query(MappingDefinition).filter(
                MappingDefinition.document_id == doc.id, MappingDefinition.status == "ACTIVE"
            ).all()
            matched = [r for r in rows if needle in str(r.target_attribute or "").lower() or needle in str(r.target_expression or "").lower()]
            if not matched:
                continue
            family = doc.mapping_family or doc.title
            timelines.setdefault(family, []).append({
                "document": self._document_dict(doc),
                "definitions": [self._mapping_dict(r, doc) for r in matched],
            })
        for family, versions in timelines.items():
            previous = None
            for version in versions:
                current = version["definitions"][0] if version["definitions"] else None
                if previous is None:
                    version["change_from_previous"] = {"status": "INITIAL", "changes": []}
                else:
                    changes = self._changes(previous, current)
                    version["change_from_previous"] = {"status": "CHANGED" if changes else "UNCHANGED", "changes": changes}
                previous = current
        return {
            "attribute": attribute,
            "mapping_family": mapping_family,
            "timelines": [{"mapping_family": k, "versions": v} for k, v in timelines.items()],
            "evidence_rule": "Every version preserves document, sheet and row provenance.",
        }

    def compare_versions(self, db: Session, mapping_family: str, from_version: str, to_version: str,
                         attribute: str | None = None) -> dict:
        docs = self._family_documents(db, mapping_family)
        from_doc = next((d for d in docs if str(d.document_version or "").lower() == from_version.lower()), None)
        to_doc = next((d for d in docs if str(d.document_version or "").lower() == to_version.lower()), None)
        if not from_doc or not to_doc:
            raise ValueError("requested mapping family/version was not found")
        before_rows = db.query(MappingDefinition).filter(MappingDefinition.document_id == from_doc.id, MappingDefinition.status == "ACTIVE").all()
        after_rows = db.query(MappingDefinition).filter(MappingDefinition.document_id == to_doc.id, MappingDefinition.status == "ACTIVE").all()
        needle = str(attribute or "").strip().lower()
        if needle:
            before_rows = [r for r in before_rows if needle in str(r.target_attribute or "").lower() or needle in str(r.target_expression or "").lower()]
            after_rows = [r for r in after_rows if needle in str(r.target_attribute or "").lower() or needle in str(r.target_expression or "").lower()]
        before = {str(r.target_expression or r.target_attribute).lower(): r for r in before_rows}
        after = {str(r.target_expression or r.target_attribute).lower(): r for r in after_rows}
        items = []
        for key in sorted(set(before) | set(after)):
            left, right = before.get(key), after.get(key)
            if left and not right:
                items.append({"status": "REMOVED", "target": left.target_expression, "from": self._mapping_dict(left, from_doc), "to": None, "changes": []})
            elif right and not left:
                items.append({"status": "ADDED", "target": right.target_expression, "from": None, "to": self._mapping_dict(right, to_doc), "changes": []})
            else:
                changes = self._changes(self._row_signature(left), self._row_signature(right))
                items.append({"status": "CHANGED" if changes else "UNCHANGED", "target": right.target_expression,
                              "from": self._mapping_dict(left, from_doc), "to": self._mapping_dict(right, to_doc), "changes": changes})
        summary = {name: sum(1 for x in items if x["status"] == name) for name in ("ADDED", "REMOVED", "CHANGED", "UNCHANGED")}
        return {
            "mapping_family": mapping_family, "from_version": from_version, "to_version": to_version,
            "attribute": attribute, "from_document": self._document_dict(from_doc), "to_document": self._document_dict(to_doc),
            "items": items, "summary": summary,
        }

    def change_impact(self, db: Session, mapping_family: str, from_version: str, to_version: str,
                      attribute: str | None = None, project: str | None = None) -> dict:
        """Return one consolidated mapping-change blast radius.

        PostgreSQL remains authoritative for mapping versions/provenance. Neo4j is
        used only for current engineering impact. Both the old and new targets are
        retained so a rename to a not-yet-existing Java field still reports the
        impact of the field being replaced.
        """
        comparison = self.compare_versions(db, mapping_family, from_version, to_version, attribute)
        validation = self.validate(db, comparison["to_document"]["id"])

        changed_items = [x for x in comparison["items"] if x["status"] != "UNCHANGED"]
        old_targets: list[str] = []
        new_targets: list[str] = []
        for item in changed_items:
            before = item.get("from") or {}
            after = item.get("to") or {}
            old = str(before.get("target_attribute") or "").strip()
            new = str(after.get("target_attribute") or "").strip()
            if old and old not in old_targets:
                old_targets.append(old)
            if new and new not in new_targets:
                new_targets.append(new)

        # A target rename is represented by compare_versions as REMOVED + ADDED.
        # Pair it explicitly when an attribute-focused comparison has one of each.
        renames = []
        removed = [x for x in changed_items if x["status"] == "REMOVED"]
        added = [x for x in changed_items if x["status"] == "ADDED"]
        if attribute and len(removed) == 1 and len(added) == 1:
            old_expr = (removed[0].get("from") or {}).get("target_expression")
            new_expr = (added[0].get("to") or {}).get("target_expression")
            if old_expr and new_expr and str(old_expr).lower() != str(new_expr).lower():
                renames.append({"from": old_expr, "to": new_expr})

        findings_by_target = {
            str(x.get("target") or "").lower(): x for x in validation.get("findings", [])
        }
        impacts = []
        try:
            from services.cross_project.engineering_knowledge_graph_service import EngineeringKnowledgeGraphService
            graph = EngineeringKnowledgeGraphService()
            for target in list(dict.fromkeys(old_targets + new_targets)):
                try:
                    impact = graph.impact(target, str(project or "").strip() or None)
                except Exception as exc:
                    impact = {"status": "UNAVAILABLE", "entity": target, "error": str(exc)}
                impacts.append({"target_attribute": target, "engineering_impact": impact})
        except Exception as exc:
            impacts.append({"target_attribute": None, "engineering_impact": {"status": "UNAVAILABLE", "error": str(exc)}})

        risks = []
        for target in new_targets:
            # Validation returns target expressions such as Student.gpaValue.
            matching = [v for k, v in findings_by_target.items() if k == target.lower() or k.endswith('.' + target.lower())]
            if matching and any(x.get("status") == "NOT_FOUND_IN_CODE_GRAPH" for x in matching):
                risks.append({
                    "severity": "HIGH",
                    "type": "TARGET_NOT_FOUND_IN_CODE",
                    "target": target,
                    "message": f"Latest mapping target {target} is not present in the current code graph.",
                })
        if comparison.get("summary", {}).get("REMOVED", 0):
            risks.append({
                "severity": "MEDIUM",
                "type": "REMOVED_MAPPING",
                "count": comparison["summary"]["REMOVED"],
                "message": "One or more mappings from the previous version are absent in the new version.",
            })

        return {
            "status": "IMPACT_FOUND" if changed_items else "NO_MAPPING_CHANGE",
            "mapping_family": mapping_family,
            "from_version": from_version,
            "to_version": to_version,
            "attribute": attribute,
            "comparison": comparison,
            "target_renames": renames,
            "latest_validation": validation,
            "old_targets": old_targets,
            "new_targets": new_targets,
            "engineering_impacts": impacts,
            "risks": risks,
            "evidence_rule": "Mapping change/provenance comes from PostgreSQL; engineering blast radius comes from the current Neo4j graph.",
        }

    def quality(self, db: Session, attribute: str | None = None, source_path: str | None = None,
                mapping_family: str | None = None) -> dict:
        """Detect conflicting/duplicate active mapping definitions with provenance."""
        self._ensure_mapping_family_column(db)
        project_path = self._project_path()
        pairs = (
            db.query(MappingDefinition, MappingDocument)
            .join(MappingDocument, MappingDocument.id == MappingDefinition.document_id)
            .filter(
                MappingDefinition.project_path == project_path,
                MappingDefinition.status == "ACTIVE",
                MappingDocument.status == "ACTIVE",
            ).all()
        )
        attr = str(attribute or "").strip().lower()
        source = str(source_path or "").strip().lower()
        family = str(mapping_family or "").strip().lower()

        # Quality compares the CURRENT definition of each mapping family, not old
        # versions inside the same family. Historical differences belong to
        # compare/history and must not be mislabeled as active conflicts.
        latest_by_family: dict[str, MappingDocument] = {}
        for _, doc in pairs:
            family_name = str(doc.mapping_family or doc.title or "").strip().lower()
            current = latest_by_family.get(family_name)
            if current is None or (self._version_key(doc.document_version), int(doc.id or 0)) > (self._version_key(current.document_version), int(current.id or 0)):
                latest_by_family[family_name] = doc
        latest_ids = {doc.id for doc in latest_by_family.values()}

        evidence = []
        for row, doc in pairs:
            if doc.id not in latest_ids:
                continue
            if family and str(doc.mapping_family or doc.title or "").strip().lower() != family:
                continue
            if attr and not (
                attr == str(row.target_attribute or "").strip().lower()
                or attr in str(row.target_expression or "").strip().lower()
                or attr == str(row.source_path or "").strip().lower().split("/")[-1].split(".")[-1]
            ):
                continue
            if source and str(row.source_path or "").strip().lower() != source:
                continue
            evidence.append(self._mapping_dict(row, doc))

        by_source: dict[str, list[dict]] = {}
        for item in evidence:
            by_source.setdefault(str(item.get("source_path") or "").strip().lower(), []).append(item)

        issues = []
        for normalized_source, items in by_source.items():
            if not normalized_source or len(items) < 2:
                continue
            targets = {str(x.get("target_expression") or "").strip().lower() for x in items}
            null_rules = {str(x.get("null_rule") or "").strip().lower() for x in items}
            transforms = {str(x.get("mapping_rule") or "").strip().lower() for x in items}
            families = {str((x.get("provenance") or {}).get("document") or "").strip().lower() for x in items}
            if len(targets) > 1:
                issues.append({"severity": "HIGH", "type": "COMPETING_TARGETS", "source_path": items[0].get("source_path"), "definitions": items})
            if len(null_rules) > 1:
                issues.append({"severity": "HIGH", "type": "CONFLICTING_NULL_RULES", "source_path": items[0].get("source_path"), "definitions": items})
            if len(transforms) > 1:
                issues.append({"severity": "MEDIUM", "type": "CONFLICTING_TRANSFORMS", "source_path": items[0].get("source_path"), "definitions": items})
            if len(targets) == 1 and len(null_rules) == 1 and len(transforms) == 1 and len(families) > 1:
                issues.append({"severity": "LOW", "type": "DUPLICATE_ACTIVE_DEFINITION", "source_path": items[0].get("source_path"), "definitions": items})

        # Validate every involved document once, so quality also catches targets
        # that are active in mapping metadata but absent from current code.
        document_ids = sorted({int((x.get("provenance") or {}).get("document_id")) for x in evidence if (x.get("provenance") or {}).get("document_id") is not None})
        missing_targets = []
        for document_id in document_ids:
            validation = self.validate(db, document_id)
            for finding in validation.get("findings", []):
                if finding.get("status") == "NOT_FOUND_IN_CODE_GRAPH":
                    missing_targets.append(finding)
        if missing_targets:
            issues.append({"severity": "HIGH", "type": "TARGET_NOT_FOUND_IN_CODE", "definitions": missing_targets})

        summary = {
            "definitions_checked": len(evidence),
            "issues": len(issues),
            "high": sum(x.get("severity") == "HIGH" for x in issues),
            "medium": sum(x.get("severity") == "MEDIUM" for x in issues),
            "low": sum(x.get("severity") == "LOW" for x in issues),
        }
        return {
            "status": "ISSUES_FOUND" if issues else "CLEAN",
            "attribute": attribute,
            "source_path": source_path,
            "mapping_family": mapping_family,
            "summary": summary,
            "issues": issues,
            "evidence": evidence,
            "evidence_rule": "Every quality finding includes authoritative mapping provenance; code-existence findings are verified against Neo4j.",
        }

    def validate(self, db: Session, document_id: int | None = None) -> dict:
        query = db.query(MappingDefinition).filter(
            MappingDefinition.project_path == self._project_path(), MappingDefinition.status == "ACTIVE"
        )
        if document_id is not None:
            query = query.filter(MappingDefinition.document_id == document_id)
        rows = query.all()
        target_names = {str(r.target_attribute or "").strip().lower() for r in rows if r.target_attribute}
        graph_attributes = self._neo4j_attribute_names()
        findings = []
        for row in rows:
            target = str(row.target_attribute or "").strip().lower()
            status = "FOUND_IN_CODE_GRAPH" if target and target in graph_attributes else "NOT_FOUND_IN_CODE_GRAPH"
            findings.append({"mapping_id": row.id, "target": row.target_expression, "status": status,
                             "evidence": {"document_id": row.document_id, "sheet": row.sheet_name, "row": row.row_number}})
        return {"checked": len(rows), "targets": sorted(target_names), "findings": findings,
                "summary": {"found": sum(x["status"] == "FOUND_IN_CODE_GRAPH" for x in findings),
                            "not_found": sum(x["status"] == "NOT_FOUND_IN_CODE_GRAPH" for x in findings)}}

    def _neo4j_attribute_names(self) -> set[str]:
        if not getattr(settings, "NEO4J_ENABLED", False):
            return set()
        driver = None
        try:
            from neo4j import GraphDatabase
            driver = GraphDatabase.driver(settings.NEO4J_URI, auth=(settings.NEO4J_USERNAME, settings.NEO4J_PASSWORD))
            with driver.session(database=settings.NEO4J_DATABASE) as session:
                result = session.run("MATCH (a:EngineeringKnowledge:Attribute) RETURN a.name AS name")
                return {str(r["name"] or "").lower() for r in result if r["name"]}
        except Exception:
            return set()
        finally:
            if driver:
                driver.close()

    @staticmethod
    def _conflicts(rows: list[dict]) -> list[dict]:
        grouped: dict[tuple[str, str], list[dict]] = {}
        for row in rows:
            key = (str(row.get("source_path") or "").lower(), str(row.get("target_expression") or "").lower())
            grouped.setdefault(key, []).append(row)
        conflicts = []
        for (source, target), items in grouped.items():
            rules = {(str(x.get("mapping_rule") or "").strip().lower(), str(x.get("null_rule") or "").strip().lower()) for x in items}
            if len(items) > 1 and len(rules) > 1:
                conflicts.append({"source": source, "target": target, "definitions": items})
        return conflicts

    def sync_document_to_neo4j(self, db: Session, document_id: int) -> dict:
        if not getattr(settings, "NEO4J_ENABLED", False):
            return {"enabled": False, "status": "DISABLED"}
        document = db.query(MappingDocument).filter(MappingDocument.id == document_id).first()
        if not document:
            raise ValueError("mapping document not found")
        rows = db.query(MappingDefinition).filter(MappingDefinition.document_id == document_id, MappingDefinition.status == "ACTIVE").all()
        project = Path(document.project_path).name or document.project_path
        from neo4j import GraphDatabase
        driver = GraphDatabase.driver(settings.NEO4J_URI, auth=(settings.NEO4J_USERNAME, settings.NEO4J_PASSWORD))
        try:
            with driver.session(database=settings.NEO4J_DATABASE) as session:
                session.run("MATCH (d:EngineeringKnowledge:MappingDocument {mapping_document_id:$id}) DETACH DELETE d", id=document.id)
                session.run("""
                    MERGE (d:EngineeringKnowledge:MappingDocument {mapping_document_id:$id})
                    SET d.name=$name,d.filename=$filename,d.mapping_family=$mapping_family,d.version=$version,d.project=$project,d.source_ref=$source_ref
                """, id=document.id, name=document.title, filename=document.filename, mapping_family=document.mapping_family or document.title, version=document.document_version,
                    project=project, source_ref=document.source_ref)
                for row in rows:
                    session.run("""
                        MATCH (d:EngineeringKnowledge:MappingDocument {mapping_document_id:$document_id})
                        MERGE (m:EngineeringKnowledge:MappingDefinition {mapping_id:$mapping_id})
                        SET m.name=$mapping_name,m.project=$project,m.source_type=$source_type,m.source_path=$source_path,
                            m.mapping_rule=$mapping_rule,m.null_rule=$null_rule,m.validation_rule=$validation_rule,
                            m.comments=$comments,m.sheet=$sheet,m.row=$row,m.target_expression=$target_expression
                        MERGE (m)-[:DEFINED_IN]->(d)
                        MERGE (s:EngineeringKnowledge:MappingSource {id:$source_id})
                        SET s.name=$source_path,s.source_type=$source_type,s.source_system=$source_system,s.project=$project
                        MERGE (s)-[:SOURCE_OF]->(m)
                        WITH s,m
                        OPTIONAL MATCH (a:EngineeringKnowledge:Attribute)
                        WHERE (toLower(coalesce(a.attribute,''))=toLower($target_attribute)
                               OR toLower(a.name)=toLower($target_attribute)
                               OR toLower(a.name)=toLower($target_expression))
                          AND (a.project=$project OR a.project IS NULL)
                          AND ($target_class IS NULL OR toLower(coalesce(a.owner,''))=toLower($target_class)
                               OR toLower(a.name)=toLower($target_expression))
                        FOREACH (_ IN CASE WHEN a IS NULL THEN [] ELSE [1] END |
                            MERGE (s)-[r:MAPS_TO]->(a)
                            SET r.mapping_id=$mapping_id,r.document=$document_name,r.sheet=$sheet,r.row=$row,
                                r.mapping_rule=$mapping_rule,r.null_rule=$null_rule,r.validation_rule=$validation_rule
                            MERGE (m)-[:TARGETS]->(a)
                        )
                    """, document_id=document.id, mapping_id=row.id, mapping_name=f"{document.filename}:{row.sheet_name}:{row.row_number}",
                        project=project, source_type=row.source_type, source_path=row.source_path, source_system=row.source_system,
                        mapping_rule=row.mapping_rule, null_rule=row.null_rule, validation_rule=row.validation_rule,
                        comments=row.comments, sheet=row.sheet_name, row=row.row_number, target_expression=row.target_expression,
                        target_attribute=row.target_attribute, target_class=row.target_class, document_name=document.filename,
                        source_id=hashlib.sha1(f"{project}|{row.source_type}|{row.source_system}|{row.source_path}".encode()).hexdigest()[:24])
            return {"enabled": True, "status": "SYNCED", "document_id": document.id, "mappings": len(rows)}
        finally:
            driver.close()

    @staticmethod
    def _document_dict(row: MappingDocument) -> dict:
        return {"id": row.id, "filename": row.filename, "title": row.title, "mapping_family": row.mapping_family or row.title, "document_version": row.document_version,
                "source_ref": row.source_ref, "project_path": row.project_path, "sheet_count": row.sheet_count,
                "row_count": row.row_count, "status": row.status, "created_at": row.created_at.isoformat() if row.created_at else None}

    @staticmethod
    def _mapping_dict(row: MappingDefinition, document: MappingDocument | None) -> dict:
        return {"id": row.id, "source_type": row.source_type, "source_system": row.source_system, "source_path": row.source_path,
                "mapping_rule": row.mapping_rule, "target_class": row.target_class, "target_attribute": row.target_attribute,
                "target_expression": row.target_expression, "null_rule": row.null_rule, "validation_rule": row.validation_rule,
                "comments": row.comments, "provenance": {"document_id": row.document_id,
                    "document": document.filename if document else None, "document_version": document.document_version if document else None,
                    "mapping_family": (document.mapping_family or document.title) if document else None,
                    "sheet": row.sheet_name, "row": row.row_number, "source_ref": document.source_ref if document else None}}
